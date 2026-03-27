from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"D:\ccgl room")
OUT_DIR = ROOT / "docs" / "slides"
OUT_DIR.mkdir(parents=True, exist_ok=True)

SCREEN_A = ROOT / "triad-live-screen.png"
SCREEN_B = ROOT / "triad-live-screen-2.png"
SCREEN_C = ROOT / "triad-workbench-preview.png"

TITLE_COLOR = RGBColor(18, 32, 51)
ACCENT = RGBColor(38, 92, 214)
TEXT = RGBColor(40, 52, 71)
MUTED = RGBColor(85, 103, 130)
BG = RGBColor(245, 248, 252)
CARD = RGBColor(255, 255, 255)
NIGHT = RGBColor(16, 24, 38)
WHITE = RGBColor(255, 255, 255)
SUCCESS = RGBColor(17, 125, 74)
WARN = RGBColor(181, 113, 0)


def make_prs() -> tuple[Presentation, object]:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs, prs.slide_layouts[6]


def add_background(slide, accent: RGBColor = ACCENT) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = accent
    band.line.fill.background()


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.38), Inches(12), Inches(0.85))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Aptos Display"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = TITLE_COLOR
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.62), Inches(1.0), Inches(12), Inches(0.4))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        r.font.name = "Aptos"
        r.font.size = Pt(13)
        r.font.color.rgb = MUTED


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12), Inches(0.22))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, font_size: int = 20) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.bullet = True
        p.space_after = Pt(8)
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT


def add_numbered(slide, items: list[str], x: float, y: float, w: float, h: float, font_size: int = 18) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items, start=1):
        p = tf.paragraphs[0] if idx == 1 else tf.add_paragraph()
        p.text = f"{idx}. {item}"
        p.space_after = Pt(8)
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT


def add_card(slide, title: str, body: list[str], x: float, y: float, w: float, h: float, title_color: RGBColor = TITLE_COLOR) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = RGBColor(221, 229, 239)

    title_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.3), Inches(0.38))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Aptos"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = title_color

    add_bullets(slide, body, x + 0.18, y + 0.5, w - 0.32, h - 0.65, font_size=15)


def add_code_block(slide, code: str, x: float, y: float, w: float, h: float) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NIGHT
    shape.line.fill.background()
    box = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.14), Inches(w - 0.24), Inches(h - 0.2))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = code
    r.font.name = "Cascadia Code"
    r.font.size = Pt(14)
    r.font.color.rgb = WHITE


def add_picture_if_exists(slide, path: Path, x: float, y: float, w: float) -> None:
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


def build_client_demo() -> Path:
    prs, blank = make_prs()
    out = OUT_DIR / "triad-workbench-client-demo.pptx"

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Triad Workbench", "Client/demo deck")
    add_bullets(
        slide,
        [
            "One desktop app for Claude, Codex, Gemini, and Ollama",
            "Local-first orchestration with git worktree safety",
            "Review, verify, and promote changes instead of blindly applying them",
        ],
        0.8,
        1.7,
        5.9,
        2.4,
    )
    add_card(slide, "Why it matters", ["Safer than direct repo editing", "Faster than juggling browser tabs", "Auditable via project archive"], 7.1, 1.65, 5.3, 2.15)
    add_picture_if_exists(slide, SCREEN_C, 0.8, 4.2, 11.8)
    add_footer(slide, "Client/demo overview")

    slide = prs.slides.add_slide(blank)
    add_background(slide, SUCCESS)
    add_title(slide, "What You See In The App", "A single mission-control screen")
    add_card(slide, "Top bar", ["Open project", "Switch runner", "Probe agents"], 0.7, 1.55, 2.4, 1.5)
    add_card(slide, "Agent grid", ["Claude", "Codex", "Gemini", "Ollama"], 0.7, 3.25, 2.4, 1.5)
    add_card(slide, "Task system", ["Workflow list", "Task history", "Review detail"], 0.7, 4.95, 2.4, 1.6)
    add_picture_if_exists(slide, SCREEN_A, 3.45, 1.45, 9.2)
    add_footer(slide, "Visual tour")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "How The Safe Coding Loop Works", "The built-in process")
    add_numbered(
        slide,
        [
            "Open a git project",
            "Run a workflow",
            "Triad creates a dedicated task worktree",
            "Agents work inside that isolated branch",
            "Review the patch, logs, and findings",
            "Promote only after approval",
        ],
        0.9,
        1.8,
        5.4,
        4.5,
        20,
    )
    add_card(slide, "Default loop", ["Claude writes", "Codex reviews", "Claude fixes", "Codex verifies"], 7.0, 1.9, 5.0, 2.2)
    add_card(slide, "Safety promise", ["No direct main-checkout edits", "Explicit promotion required", "Project archive keeps evidence"], 7.0, 4.35, 5.0, 1.85, title_color=SUCCESS)
    add_footer(slide, "Workflow model")

    slide = prs.slides.add_slide(blank)
    add_background(slide, WARN)
    add_title(slide, "Agent Roles At A Glance", "Use each model for the right job")
    add_card(slide, "Claude", ["Primary implementation", "Best for writing code and applying fixes"], 0.7, 1.7, 2.8, 1.8)
    add_card(slide, "Codex", ["Review and verification", "Great for bug-finding and test-minded critique"], 3.7, 1.7, 2.8, 1.8)
    add_card(slide, "Gemini", ["Architecture and second opinion", "Useful in compare workflows"], 6.7, 1.7, 2.8, 1.8)
    add_card(slide, "Ollama", ["Local support model", "Guide / Tester / Developer / Monitor"], 9.7, 1.7, 2.8, 1.8)
    add_bullets(
        slide,
        [
            "Direct terminals are available for CLI agents",
            "Workflows are the main product experience",
            "Ollama is strongest as local compare or monitor support",
        ],
        0.9,
        4.2,
        11.7,
        1.6,
        19,
    )
    add_footer(slide, "Agent positioning")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Using It On An Existing Repo", "The normal daily flow")
    add_picture_if_exists(slide, SCREEN_B, 7.0, 1.45, 5.6)
    add_numbered(
        slide,
        [
            "Open the repository",
            "Connect the agents you need",
            "Pick a workflow",
            "Write a small, clear task brief",
            "Review findings before promotion",
        ],
        0.85,
        1.8,
        5.6,
        3.8,
        20,
    )
    add_footer(slide, "Existing project use")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Starting From Scratch", "A simple new-project path")
    add_code_block(slide, "mkdir my-small-app\ncd my-small-app\ngit init\nnpm init -y\ngit add .\ngit commit -m \"Initial scaffold\"", 0.85, 1.8, 5.8, 2.0)
    add_card(slide, "Then in Triad", ["Open the folder", "Connect Claude and Codex", "Run a tiny first feature with tests"], 7.0, 1.85, 5.0, 1.85)
    add_card(slide, "Best first task", ["Create one feature", "Add tests", "Keep scope narrow", "Inspect before applying"], 7.0, 4.1, 5.0, 1.75, title_color=SUCCESS)
    add_footer(slide, "New project use")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Review And Promotion", "Where the app becomes safer than normal agent use")
    add_card(slide, "Review center", ["Summary", "Prompt", "Stdout / stderr", "Patch", "Findings"], 0.8, 1.8, 3.4, 2.1)
    add_card(slide, "Promotion actions", ["Apply to main", "Keep worktree", "Open task branch"], 4.45, 1.8, 3.2, 2.1)
    add_card(slide, "Archive", ["Stores snapshots, artifacts, prompts, logs, transcripts, and events"], 7.9, 1.8, 4.5, 2.1)
    add_bullets(slide, ["This makes the app suitable for controlled team-style coding, not just solo prompting."], 0.9, 4.7, 11.5, 1.0, 21)
    add_footer(slide, "Safety layer")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Best Current Use Cases", "What to demo or start with first")
    add_bullets(
        slide,
        [
            "Small to medium feature work",
            "Safe code review loops",
            "Architecture comparison before implementation",
            "Projects where you want an auditable AI trail",
        ],
        0.85,
        1.8,
        6.2,
        3.0,
        22,
    )
    add_card(slide, "Not the main v0.1 focus", ["Huge one-shot app generation", "Per-agent full chat product", "Blind auto-merge without review"], 7.35, 1.95, 4.8, 2.0, title_color=WARN)
    add_footer(slide, "Demo close")

    prs.save(str(out))
    return out


def build_beginner() -> Path:
    prs, blank = make_prs()
    out = OUT_DIR / "triad-workbench-beginner-5-slide.pptx"

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Triad Workbench", "5-slide beginner guide")
    add_bullets(slide, ["A desktop app that helps multiple coding agents work safely on one repo.", "Best first use: small git project + Code -> Review -> Fix -> Verify."], 0.9, 1.9, 11.4, 2.2, 24)
    add_footer(slide, "Beginner guide")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Step 1: Open A Project", "Start with a real git repo")
    add_numbered(slide, ["Click Open project", "Choose a git-backed folder", "Probe agents", "Connect Claude and Codex first"], 0.9, 1.8, 5.2, 3.5, 21)
    add_picture_if_exists(slide, SCREEN_A, 6.4, 1.55, 6.0)
    add_footer(slide, "Step 1")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Step 2: Pick A Workflow", "Use the safe default first")
    add_card(slide, "Recommended", ["Code -> Review -> Fix -> Verify"], 0.9, 1.8, 4.0, 1.25, title_color=SUCCESS)
    add_bullets(slide, ["Claude writes", "Codex reviews", "Claude fixes", "Codex verifies"], 0.95, 3.4, 4.0, 2.0, 21)
    add_card(slide, "Write a small brief", ["One feature", "One bug", "One page", "Ask for tests"], 6.2, 1.9, 5.3, 2.0)
    add_footer(slide, "Step 2")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Step 3: Review The Result", "Never skip the review center")
    add_bullets(slide, ["Check findings", "Inspect the patch", "Look at stdout / stderr if something failed", "Only promote after review"], 0.9, 1.8, 5.8, 3.2, 22)
    add_card(slide, "Promotion buttons", ["Apply to main", "Keep worktree", "Open task branch"], 7.05, 2.0, 4.7, 1.8)
    add_footer(slide, "Step 3")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Step 4: Try This First Prompt", "A safe beginner task")
    add_code_block(slide, "Create a TypeScript CLI todo app with add, list, and done commands.\nStore data in a local JSON file.\nAdd Vitest tests for the main flows.\nKeep the implementation simple and readable.", 0.95, 1.9, 11.2, 2.2)
    add_bullets(slide, ["Small tasks work best", "Keep your main branch clean", "Use Save now if you want an archive snapshot"], 1.0, 4.8, 10.8, 1.3, 20)
    add_footer(slide, "First real run")

    prs.save(str(out))
    return out


def build_team_training() -> Path:
    prs, blank = make_prs()
    out = OUT_DIR / "triad-workbench-team-training.pptx"

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Triad Workbench Team Training", "How to use the app consistently across a team")
    add_bullets(slide, ["Common workflow language", "Safer review discipline", "Example prompts and operating rules"], 0.9, 1.9, 11.0, 2.0, 24)
    add_footer(slide, "Team training")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Operating Model", "What the team should expect from Triad")
    add_bullets(
        slide,
        [
            "Triad is workflow-first, not chat-first",
            "Each task gets its own worktree",
            "Artifacts and findings are part of the output, not optional extras",
            "Promotion is a human decision point",
        ],
        0.9,
        1.8,
        11.6,
        3.2,
        21,
    )
    add_picture_if_exists(slide, SCREEN_C, 0.95, 4.8, 11.3)
    add_footer(slide, "Team model")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Who To Use For What", "Role guidance for the team")
    add_card(slide, "Claude", ["Implementation", "Fixes", "Main coding pass"], 0.8, 1.8, 2.8, 1.9)
    add_card(slide, "Codex", ["Review", "Verification", "Bug-finding"], 3.85, 1.8, 2.8, 1.9)
    add_card(slide, "Gemini", ["Architecture compare", "Extra critique", "Low-cost second opinion"], 6.9, 1.8, 2.8, 1.9)
    add_card(slide, "Ollama", ["Local monitoring", "Local compare", "Support model"], 9.95, 1.8, 2.6, 1.9)
    add_footer(slide, "Agent roles")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Standard Team Workflow", "The recommended default")
    add_numbered(
        slide,
        [
            "Open clean git repo",
            "Connect required agents",
            "Run Code -> Review -> Fix -> Verify",
            "Inspect findings and patch",
            "Promote only after acceptance",
        ],
        0.9,
        1.85,
        5.2,
        3.6,
        21,
    )
    add_card(slide, "Use this for", ["Feature work", "Bug fixes", "Small refactors", "Safe iterative changes"], 7.0, 1.95, 5.0, 2.1)
    add_footer(slide, "Default workflow")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Team Rules", "These reduce bad runs and messy diffs")
    add_bullets(
        slide,
        [
            "One narrow task per run",
            "Always ask for tests when code changes",
            "Do not ask for a full app build in one prompt",
            "Keep main checkout clean before Apply to main",
            "Use Keep worktree when the result needs manual refinement",
        ],
        0.9,
        1.8,
        11.4,
        3.6,
        22,
    )
    add_footer(slide, "Team rules")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Example Prompt: Small Feature", "Ready to paste into the task brief box")
    add_code_block(slide, "Add a small React settings panel for notification preferences.\nInclude a save button and simple local validation.\nKeep the UI responsive.\nAdd tests for render and save behavior.\nDo not modify unrelated files.", 0.9, 1.9, 11.2, 2.4)
    add_bullets(slide, ["Why this is good: narrow scope, stack named, tests required, unrelated files protected."], 1.0, 4.8, 11.0, 1.0, 19)
    add_footer(slide, "Prompt example 1")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Example Prompt: API Task", "Another good team-style prompt")
    add_code_block(slide, "Create an Express route for creating and listing notes.\nUse in-memory storage first.\nReturn clean JSON responses.\nAdd supertest coverage for create and list.\nKeep the implementation simple and readable.", 0.9, 1.9, 11.2, 2.4)
    add_bullets(slide, ["Why this is good: one bounded backend slice, simple storage, explicit test framework."], 1.0, 4.8, 11.0, 1.0, 19)
    add_footer(slide, "Prompt example 2")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Example Prompt: Architecture Compare", "Use before implementation when the design is unclear")
    add_code_block(slide, "Compare three approaches for adding plugin support to this app:\n1. local config-driven plugins\n2. IPC-managed extension loader\n3. MCP-backed extension model\nEvaluate complexity, safety, and maintainability.\nDo not write code; produce findings and recommendation only.", 0.9, 1.8, 11.2, 2.5)
    add_bullets(slide, ["Use the Architecture Compare workflow for prompts like this."], 1.0, 4.8, 8.0, 1.0, 19)
    add_footer(slide, "Prompt example 3")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "What Reviewers Should Check", "How to read the output")
    add_card(slide, "Findings", ["Severity", "Evidence", "Recommended action"], 0.8, 1.9, 3.0, 1.7)
    add_card(slide, "Artifacts", ["Prompt", "Stdout", "Stderr", "Patch"], 4.05, 1.9, 3.0, 1.7)
    add_card(slide, "Promotion", ["Apply to main", "Keep worktree", "Open task branch"], 7.3, 1.9, 5.0, 1.7)
    add_picture_if_exists(slide, SCREEN_B, 0.95, 4.1, 11.2)
    add_footer(slide, "Review habits")

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Common Mistakes To Avoid", "What to coach during rollout")
    add_bullets(
        slide,
        [
            "Running on a non-git folder",
            "Using a giant vague prompt",
            "Skipping Probe agents after changing runner",
            "Treating Ollama like a full direct-chat pane",
            "Applying to main without reading findings",
        ],
        0.95,
        1.9,
        11.0,
        3.3,
        22,
    )
    add_footer(slide, "Pitfalls")

    prs.save(str(out))
    return out


if __name__ == "__main__":
    created = [build_client_demo(), build_beginner(), build_team_training()]
    for path in created:
        print(f"Created {path}")
