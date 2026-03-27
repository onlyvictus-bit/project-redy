from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"D:\ccgl room")
OUT_DIR = ROOT / "docs" / "slides"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_FILE = OUT_DIR / "triad-workbench-user-guide.pptx"

SCREENSHOT = ROOT / "triad-live-screen.png"

TITLE_COLOR = RGBColor(18, 32, 51)
ACCENT = RGBColor(38, 92, 214)
TEXT = RGBColor(40, 52, 71)
MUTED = RGBColor(85, 103, 130)
BG = RGBColor(245, 248, 252)
CARD = RGBColor(255, 255, 255)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.62), Inches(1.08), Inches(11.5), Inches(0.45))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        run.font.name = "Aptos"
        run.font.size = Pt(13)
        run.font.color.rgb = MUTED


def add_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, font_size: int = 20) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.bullet = True


def add_numbered(slide, items: list[str], x: float, y: float, w: float, h: float, font_size: int = 18) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items, start=1):
        p = tf.paragraphs[0] if i == 1 else tf.add_paragraph()
        p.text = f"{i}. {item}"
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(10)


def add_card(slide, title: str, body: list[str], x: float, y: float, w: float, h: float) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = RGBColor(221, 229, 239)
    title_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.3), Inches(0.4))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Aptos"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR
    add_bullets(slide, body, x + 0.18, y + 0.55, w - 0.35, h - 0.7, font_size=15)


def add_code_block(slide, code: str, x: float, y: float, w: float, h: float) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(16, 24, 38)
    shape.line.fill.background()
    box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.14), Inches(w - 0.25), Inches(h - 0.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = code
    run.font.name = "Cascadia Code"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(235, 241, 248)


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "CCGL / Triad Workbench", "Multi-agent coding workbench guide")
add_bullets(
    slide,
    [
        "Understand what the app does and how the UI is organized",
        "Learn each workflow, agent option, and promotion action",
        "See the safe path for both new projects and existing repos",
    ],
    0.8,
    1.8,
    6.0,
    2.0,
)
add_card(
    slide,
    "Core agents",
    ["Claude Code", "Codex CLI", "Gemini CLI", "Ollama"],
    7.1,
    1.7,
    2.35,
    2.0,
)
add_card(
    slide,
    "Core promise",
    ["Local-first", "CLI-native", "Worktree-safe", "Review before promotion"],
    9.7,
    1.7,
    2.7,
    2.0,
)
add_footer(slide, "Triad Workbench user guide")


slides_data = [
    (
        "What Triad Workbench Is",
        "A local mission-control app for coding agents",
        [
            "Electron desktop app that orchestrates installed coding CLIs",
            "Built around isolated git worktrees rather than direct main-branch edits",
            "Best fit: controlled implementation, review, verification, and promotion",
            "Primary loop: Claude writes -> Codex reviews -> Claude fixes -> Codex verifies",
        ],
    ),
    (
        "Why Use It",
        "Safer than asking one AI to edit your repo directly",
        [
            "Keeps each task in its own worktree",
            "Preserves prompts, logs, diffs, and findings in a project archive",
            "Lets multiple agents critique the same task from different roles",
            "Requires explicit promotion before code reaches your main checkout",
        ],
    ),
    (
        "Core Concepts",
        "The mental model to keep in mind while using the app",
        [
            "Agent cards = current tool status, role, terminal, and quick actions",
            "Workflow = built-in chain of agent steps",
            "Task = one run with its own worktree, artifacts, findings, and stage history",
            "Review center = inspect prompts, logs, patches, and findings before promoting",
            "Archive = per-project evidence folder at .triad-workbench",
        ],
    ),
    (
        "Main Screen Tour",
        "What each area of the UI is for",
        None,
    ),
]

for title, subtitle, bullets in slides_data:
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, title, subtitle)
    if bullets:
        add_bullets(slide, bullets, 0.8, 1.7, 11.8, 4.8)
    else:
        add_card(slide, "Top bar", ["Open project", "Runner selector", "Probe agents"], 0.75, 1.55, 2.8, 1.6)
        add_card(slide, "Left rail", ["Workflow picker", "Task list", "Run history"], 0.75, 3.35, 2.8, 1.7)
        add_card(slide, "Center grid", ["Claude", "Codex", "Gemini", "Ollama"], 0.75, 5.25, 2.8, 1.55)
        add_card(slide, "Right rail", ["Findings", "Ollama model", "Archive", "Notifications"], 3.8, 1.55, 2.5, 1.9)
        add_card(slide, "Bottom composer", ["Write the task brief", "Pick a workflow", "Run the task"], 3.8, 3.75, 2.5, 1.7)
        if SCREENSHOT.exists():
            slide.shapes.add_picture(str(SCREENSHOT), Inches(6.6), Inches(1.5), width=Inches(6.0))
        add_footer(slide, "Screen tour")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Agents And What To Use Them For", "Each card has a different job in the system")
add_card(slide, "Claude Code", ["Primary coder", "Best for implementing features", "Also used for fix passes"], 0.7, 1.55, 2.9, 1.9)
add_card(slide, "Codex CLI", ["Reviewer / tester", "Best for bug finding and verification", "Used in final review steps"], 3.85, 1.55, 2.9, 1.9)
add_card(slide, "Gemini CLI", ["Architecture or extra critique", "Good second opinion", "Useful in compare workflows"], 7.0, 1.55, 2.9, 1.9)
add_card(slide, "Ollama", ["Local model", "Guide / Tester / Developer / Monitor personalities", "Great for local compare or away-monitor"], 10.15, 1.55, 2.5, 1.9)
add_bullets(
    slide,
    [
        "Connect = log in or authorize a CLI agent",
        "Open terminal = direct interactive session with that CLI agent",
        "Role / personality changes how the agent is used in workflows",
        "Ready is required before a workflow can use that agent",
    ],
    0.9,
    4.0,
    11.7,
    2.1,
    18,
)
add_footer(slide, "Agent options")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Built-In Workflows", "Pick the workflow that matches the job")
add_card(slide, "Code -> Review -> Fix -> Verify", ["Default and safest", "Claude codes", "Codex reviews", "Claude fixes", "Codex verifies"], 0.7, 1.55, 2.85, 2.35)
add_card(slide, "Code -> Gemini Compare -> Codex Review", ["Adds Gemini as an extra critic", "Good for architecture-sensitive changes"], 3.75, 1.55, 2.85, 2.35)
add_card(slide, "Architecture Compare", ["Runs side-by-side design critique", "Claude + Codex + Gemini + Ollama"], 6.8, 1.55, 2.85, 2.35)
add_card(slide, "Away Monitor", ["Uses Ollama as a local monitor", "Good for long-running sessions or local observation"], 9.85, 1.55, 2.75, 2.35)
add_bullets(
    slide,
    [
        "Use Code -> Review -> Fix -> Verify for your first real task",
        "Use Architecture Compare before coding when the design is still unclear",
        "Use Away Monitor when you want Ollama watching a longer task flow",
    ],
    0.85,
    4.45,
    11.8,
    1.8,
    18,
)
add_footer(slide, "Workflow guide")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "How A Task Works", "What happens after you click Run workflow")
add_numbered(
    slide,
    [
        "Triad inspects the selected git project",
        "A dedicated task worktree is created",
        "The selected workflow runs one agent step at a time",
        "Each step produces artifacts: prompt, stdout, stderr, summary, findings, patch",
        "The review center shows the results",
        "You explicitly promote, keep, or open the task worktree",
    ],
    0.9,
    1.7,
    11.6,
    4.8,
    19,
)
add_footer(slide, "Task lifecycle")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Using Triad On An Existing Project", "The normal way to use the app")
add_numbered(
    slide,
    [
        "Open a real git repository with Open project",
        "Set the correct runner (Windows or WSL)",
        "Probe agents and connect any that are missing login",
        "Choose a workflow",
        "Write a small, clear task brief",
        "Run the workflow",
        "Review findings, logs, and patch output",
        "Apply to main, keep worktree, or open the task branch",
    ],
    0.8,
    1.7,
    6.0,
    5.2,
    18,
)
add_card(
    slide,
    "Important rules",
    [
        "Project must be a git repo",
        "Main checkout should be clean before Apply to main",
        "Keep tasks small and reviewable",
        "Use Save now if you want an archive snapshot immediately",
    ],
    7.2,
    1.8,
    5.1,
    3.0,
)
add_footer(slide, "Existing project flow")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Starting A New Project From Scratch", "Minimal safe path for a fresh repo")
add_code_block(
    slide,
    "mkdir my-small-app\ncd my-small-app\ngit init\nnpm init -y\ngit add .\ngit commit -m \"Initial scaffold\"",
    0.8,
    1.7,
    5.7,
    2.1,
)
add_numbered(
    slide,
    [
        "Create the folder and initialize git",
        "Add a tiny starter file or package manifest",
        "Make the first commit before asking Triad to change anything",
        "Open the folder in Triad",
        "Start with one small task, not a full app request",
    ],
    0.85,
    4.2,
    5.8,
    2.2,
    17,
)
add_card(
    slide,
    "Recommended first task",
    [
        "Create a TypeScript CLI todo app with add, list, and done commands.",
        "Store data in a local JSON file.",
        "Add Vitest tests for the main flows.",
        "Keep the implementation simple and readable.",
    ],
    6.9,
    1.8,
    5.55,
    3.7,
)
add_footer(slide, "New project quickstart")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "How To Write Good Task Briefs", "Small, explicit prompts work best")
add_bullets(
    slide,
    [
        "Say exactly what to build",
        "Mention the language or framework",
        "State constraints: keep simple, avoid unrelated files, add tests",
        "Ask for the shape of success, not vague creativity",
    ],
    0.85,
    1.7,
    5.7,
    2.5,
    20,
)
add_code_block(
    slide,
    "Build a small Express API with routes for create, list, and delete notes.\nUse in-memory storage first.\nAdd tests with supertest.\nDo not modify unrelated files.",
    6.8,
    1.8,
    5.7,
    2.6,
)
add_card(
    slide,
    "Avoid briefs like",
    [
        "Build the whole app from scratch",
        "Refactor everything",
        "Make it production ready in one step",
    ],
    6.8,
    4.8,
    5.7,
    1.45,
)
add_footer(slide, "Prompting guide")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Review Center, Promotion, And Archive", "The safety layer of the app")
add_card(slide, "Review center", ["Inspect summary, prompt, stdout, stderr, patch, findings, and command runs"], 0.8, 1.7, 3.8, 1.95)
add_card(slide, "Promotion actions", ["Apply to main", "Keep worktree", "Open task branch"], 4.8, 1.7, 2.8, 1.95)
add_card(slide, "Archive", ["Save prompts, logs, artifacts, transcripts, and events to .triad-workbench"], 7.8, 1.7, 4.7, 1.95)
add_bullets(
    slide,
    [
        "Apply to main is the final merge-back step",
        "Keep worktree preserves the isolated branch for more manual work",
        "Open task branch opens the task worktree folder",
        "Archive helps you audit what every agent actually did",
    ],
    0.85,
    4.1,
    11.6,
    2.0,
    18,
)
add_footer(slide, "Safety and evidence")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "What Works Best Today", "Recommended operating style for v0.1")
add_bullets(
    slide,
    [
        "Use Claude + Codex first for your main build loop",
        "Use Gemini for extra critique or architecture compare",
        "Use Ollama as local support, not as your only source of truth",
        "Start with small to medium tasks",
        "Review before promotion every time",
        "Keep your main checkout clean",
    ],
    0.85,
    1.75,
    11.8,
    4.6,
    20,
)
add_footer(slide, "Best practices")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Current Limitations And Troubleshooting", "Things to know before daily use")
add_card(
    slide,
    "Current limitations",
    [
        "Not a full per-agent chat product yet",
        "Ollama is workflow-driven rather than full direct chat",
        "Runner choice matters for CLI availability",
        "Large one-shot tasks are less reliable than smaller steps",
    ],
    0.75,
    1.7,
    5.9,
    2.5,
)
add_card(
    slide,
    "Common fixes",
    [
        "Use Probe agents after changing runner",
        "Reconnect agents if auth expires",
        "Use Open terminal for direct CLI interaction",
        "If Apply to main is blocked, clean the main checkout first",
    ],
    6.85,
    1.7,
    5.7,
    2.5,
)
add_bullets(
    slide,
    [
        "If a workflow is disabled, check project git status and required agent readiness",
        "If a connector behaves differently in WSL vs Windows, switch runner and probe again",
        "If you want direct chat with every model, that is a future feature rather than the main v0.1 workflow",
    ],
    0.85,
    4.65,
    11.8,
    1.7,
    17,
)
add_footer(slide, "Limitations")


slide = prs.slides.add_slide(blank)
add_background(slide)
add_title(slide, "Recommended First Run", "The easiest way to learn the app")
add_numbered(
    slide,
    [
        "Open a small git repo",
        "Connect Claude and Codex",
        "Choose Code -> Review -> Fix -> Verify",
        "Ask for one tiny feature plus tests",
        "Inspect findings and patch",
        "Apply to main only after review",
    ],
    0.95,
    2.0,
    11.2,
    3.8,
    21,
)
add_footer(slide, "Start here")


prs.save(str(OUT_FILE))
print(f"Created {OUT_FILE}")
